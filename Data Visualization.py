# ============================================================
# WEEK 2 TASK
# Advanced Data Visualization and Storytelling with Python
# Dataset: Titanic
# ============================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

# ------------------------------------------------------------
# 1. SETUP
# ------------------------------------------------------------

sns.set_theme(style="whitegrid")

OUTPUT_DIR = "titanic_week2_output"

if not os.path.exists(OUTPUT_DIR):
    os.makedirs(OUTPUT_DIR)

print("Loading Titanic dataset...")

# Publicly accessible Titanic dataset
url = "https://raw.githubusercontent.com/datasciencedojo/datasets/master/titanic.csv"

df = pd.read_csv(url)

print("\nOriginal Dataset:")
print(df.head())

print("\nDataset Shape:")
print(df.shape)

# ------------------------------------------------------------
# 2. DATA CLEANING
# ------------------------------------------------------------

# Standardize column names
df.columns = [col.lower().replace(" ", "_") for col in df.columns]

# Display missing values
print("\nMissing Values:")
print(df.isnull().sum())

# Fill missing Age with median
df["age"] = df["age"].fillna(df["age"].median())

# Fill missing Embarked if available
if "embarked" in df.columns:
    df["embarked"] = df["embarked"].fillna(df["embarked"].mode()[0])

# Fill missing Fare
if "fare" in df.columns:
    df["fare"] = df["fare"].fillna(df["fare"].median())

# Convert categorical values
if "sex" in df.columns:
    df["sex"] = df["sex"].str.title()

# Create useful variables
df["survival_status"] = df["survived"].map({
    0: "Did Not Survive",
    1: "Survived"
})

df["age_group"] = pd.cut(
    df["age"],
    bins=[0, 12, 18, 35, 60, 100],
    labels=[
        "Child",
        "Teenager",
        "Young Adult",
        "Adult",
        "Senior"
    ]
)

# ------------------------------------------------------------
# 3. BASIC STATISTICS
# ------------------------------------------------------------

total_passengers = len(df)
survivors = df["survived"].sum()
non_survivors = total_passengers - survivors
survival_rate = df["survived"].mean() * 100

print("\nTotal Passengers:", total_passengers)
print("Survivors:", survivors)
print("Non-Survivors:", non_survivors)
print("Overall Survival Rate:", round(survival_rate, 2), "%")

# ------------------------------------------------------------
# 4. VISUALIZATION 1
# Survival Rate by Passenger Class
# ------------------------------------------------------------

plt.figure(figsize=(9, 6))

class_survival = (
    df.groupby("pclass")["survived"]
    .mean()
    .reset_index()
)

class_survival["survived"] *= 100

ax = sns.barplot(
    data=class_survival,
    x="pclass",
    y="survived"
)

plt.title(
    "Survival Rate by Passenger Class",
    fontsize=16,
    fontweight="bold"
)

plt.xlabel("Passenger Class")
plt.ylabel("Survival Rate (%)")

for i, value in enumerate(class_survival["survived"]):
    ax.text(
        i,
        value + 2,
        f"{value:.1f}%",
        ha="center",
        fontweight="bold"
    )

plt.ylim(0, 100)
plt.tight_layout()

chart1 = os.path.join(
    OUTPUT_DIR,
    "01_survival_by_class.png"
)

plt.savefig(chart1, dpi=300)
plt.close()

# ------------------------------------------------------------
# 5. VISUALIZATION 2
# Survival Rate by Gender and Class
# ------------------------------------------------------------

plt.figure(figsize=(10, 6))

gender_class = (
    df.groupby(["pclass", "sex"])["survived"]
    .mean()
    .reset_index()
)

gender_class["survived"] *= 100

ax = sns.barplot(
    data=gender_class,
    x="pclass",
    y="survived",
    hue="sex"
)

plt.title(
    "Survival Rate by Passenger Class and Gender",
    fontsize=16,
    fontweight="bold"
)

plt.xlabel("Passenger Class")
plt.ylabel("Survival Rate (%)")

for container in ax.containers:
    ax.bar_label(
        container,
        fmt="%.1f%%",
        padding=3
    )

plt.ylim(0, 110)
plt.tight_layout()

chart2 = os.path.join(
    OUTPUT_DIR,
    "02_class_gender_survival.png"
)

plt.savefig(chart2, dpi=300)
plt.close()

# ------------------------------------------------------------
# 6. VISUALIZATION 3
# Age Distribution by Survival
# ------------------------------------------------------------

plt.figure(figsize=(10, 6))

sns.histplot(
    data=df,
    x="age",
    hue="survival_status",
    bins=30,
    kde=True,
    element="step"
)

plt.title(
    "Age Distribution of Survivors and Non-Survivors",
    fontsize=16,
    fontweight="bold"
)

plt.xlabel("Age")
plt.ylabel("Number of Passengers")

plt.tight_layout()

chart3 = os.path.join(
    OUTPUT_DIR,
    "03_age_distribution.png"
)

plt.savefig(chart3, dpi=300)
plt.close()

# ------------------------------------------------------------
# 7. VISUALIZATION 4
# Survival Rate by Age Group
# ------------------------------------------------------------

plt.figure(figsize=(10, 6))

age_survival = (
    df.groupby("age_group", observed=False)["survived"]
    .mean()
    .reset_index()
)

age_survival["survived"] *= 100

ax = sns.barplot(
    data=age_survival,
    x="age_group",
    y="survived"
)

plt.title(
    "Survival Rate Across Age Groups",
    fontsize=16,
    fontweight="bold"
)

plt.xlabel("Age Group")
plt.ylabel("Survival Rate (%)")

for container in ax.containers:
    ax.bar_label(
        container,
        fmt="%.1f%%",
        padding=3
    )

plt.ylim(0, 100)
plt.xticks(rotation=20)

plt.tight_layout()

chart4 = os.path.join(
    OUTPUT_DIR,
    "04_age_group_survival.png"
)

plt.savefig(chart4, dpi=300)
plt.close()

# ------------------------------------------------------------
# 8. VISUALIZATION 5
# Fare vs Age Scatter Plot
# ------------------------------------------------------------

plt.figure(figsize=(10, 6))

sns.scatterplot(
    data=df,
    x="age",
    y="fare",
    hue="survival_status",
    size="pclass",
    sizes=(30, 180),
    alpha=0.7
)

plt.title(
    "Relationship Between Age, Fare and Survival",
    fontsize=16,
    fontweight="bold"
)

plt.xlabel("Age")
plt.ylabel("Fare")

plt.tight_layout()

chart5 = os.path.join(
    OUTPUT_DIR,
    "05_age_fare_survival.png"
)

plt.savefig(chart5, dpi=300)
plt.close()

# ------------------------------------------------------------
# 9. VISUALIZATION 6
# Correlation Heatmap
# ------------------------------------------------------------

plt.figure(figsize=(10, 7))

numeric_columns = [
    "survived",
    "pclass",
    "age",
    "sibsp",
    "parch",
    "fare"
]

correlation = df[numeric_columns].corr()

sns.heatmap(
    correlation,
    annot=True,
    cmap="coolwarm",
    fmt=".2f",
    linewidths=0.5
)

plt.title(
    "Correlation Between Numerical Variables",
    fontsize=16,
    fontweight="bold"
)

plt.tight_layout()

chart6 = os.path.join(
    OUTPUT_DIR,
    "06_correlation_heatmap.png"
)

plt.savefig(chart6, dpi=300)
plt.close()

# ------------------------------------------------------------
# 10. ADD ANNOTATION CHART
# Overall Survival
# ------------------------------------------------------------

plt.figure(figsize=(8, 6))

counts = df["survival_status"].value_counts()

plt.pie(
    counts,
    labels=counts.index,
    autopct="%1.1f%%",
    startangle=90
)

plt.title(
    "Overall Titanic Passenger Survival",
    fontsize=16,
    fontweight="bold"
)

plt.tight_layout()

chart7 = os.path.join(
    OUTPUT_DIR,
    "07_overall_survival.png"
)

plt.savefig(chart7, dpi=300)
plt.close()

# ------------------------------------------------------------
# 11. CREATE WORD DOCUMENT
# ------------------------------------------------------------

document = Document()

# Page margins
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# ------------------------------------------------------------
# TITLE PAGE
# ------------------------------------------------------------

title = document.add_heading(
    "Advanced Data Visualization and Storytelling with Python",
    0
)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "\nTitanic Survival Analysis\n\n"
    "Week 2 Data Science Task\n\n"
    "Dataset: Titanic Passenger Dataset"
)

run.bold = True
run.font.size = Pt(16)

document.add_paragraph("\n")

intro = document.add_paragraph()

intro.add_run(
    "Objective: "
).bold = True

intro.add_run(
    "To develop advanced visualizations that transform a "
    "historical passenger dataset into a clear and meaningful "
    "data story for a non-technical audience."
)

document.add_page_break()

# ------------------------------------------------------------
# INTRODUCTION
# ------------------------------------------------------------

document.add_heading("1. Introduction", level=1)

document.add_paragraph(
    "The Titanic dataset is a well-known dataset used in "
    "data analysis and machine learning. It contains information "
    "about passengers aboard the RMS Titanic, including passenger "
    "class, gender, age, fare, family relationships and survival status."
)

document.add_paragraph(
    "The purpose of this analysis is not simply to display charts, "
    "but to use visualization as a storytelling technique. "
    "Different visualizations are used to examine passenger class, "
    "gender, age, fare and relationships between numerical variables."
)

# ------------------------------------------------------------
# DATASET
# ------------------------------------------------------------

document.add_heading("2. Dataset Description", level=1)

document.add_paragraph(
    f"The dataset contains {total_passengers} passenger records. "
    f"The main target variable is survival status, where 0 represents "
    f"non-survival and 1 represents survival."
)

table = document.add_table(
    rows=1,
    cols=3
)

table.style = "Table Grid"

hdr = table.rows[0].cells

hdr[0].text = "Variable"
hdr[1].text = "Type"
hdr[2].text = "Description"

dataset_info = [
    ("survived", "Numerical/Binary", "Passenger survival status"),
    ("pclass", "Categorical", "Passenger class"),
    ("sex", "Categorical", "Passenger gender"),
    ("age", "Numerical", "Passenger age"),
    ("fare", "Numerical", "Ticket fare"),
    ("sibsp", "Numerical", "Number of siblings/spouses"),
    ("parch", "Numerical", "Number of parents/children"),
]

for row in dataset_info:
    cells = table.add_row().cells

    for i, value in enumerate(row):
        cells[i].text = value

# ------------------------------------------------------------
# METHODS
# ------------------------------------------------------------

document.add_heading("3. Methodology", level=1)

methods = [
    "Load the publicly accessible Titanic dataset.",
    "Inspect the structure and missing values.",
    "Handle missing numerical and categorical values.",
    "Create additional variables such as age groups.",
    "Calculate descriptive statistics and survival rates.",
    "Develop multiple visualization techniques.",
    "Add captions and explanations to every visualization.",
    "Interpret the patterns from a non-technical perspective.",
    "Discuss possible business and scientific implications."
]

for method in methods:
    document.add_paragraph(
        method,
        style="List Bullet"
    )

# ------------------------------------------------------------
# KEY STATISTICS
# ------------------------------------------------------------

document.add_heading("4. Key Statistics", level=1)

document.add_paragraph(
    f"Total passengers analyzed: {total_passengers}"
)

document.add_paragraph(
    f"Passengers who survived: {survivors}"
)

document.add_paragraph(
    f"Passengers who did not survive: {non_survivors}"
)

document.add_paragraph(
    f"Overall survival rate: {survival_rate:.2f}%"
)

# ------------------------------------------------------------
# VISUALIZATION 1
# ------------------------------------------------------------

document.add_heading(
    "5. Visualization 1 – Survival Rate by Passenger Class",
    level=1
)

document.add_picture(
    chart1,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 1: Percentage of passengers who survived within each "
    "passenger class."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "The visualization compares survival rates across first, second "
    "and third class. It shows a clear difference between passenger "
    "classes. This suggests that passenger class was associated with "
    "different survival outcomes."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "A bar chart allows a non-technical audience to compare percentages "
    "quickly. The labels on top of the bars make the numerical values "
    "easy to understand without requiring the reader to inspect the axis."
)

# ------------------------------------------------------------
# VISUALIZATION 2
# ------------------------------------------------------------

document.add_heading(
    "6. Visualization 2 – Passenger Class and Gender",
    level=1
)

document.add_picture(
    chart2,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 2: Survival rates separated by passenger class and gender."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "This visualization adds gender to the passenger-class analysis. "
    "The comparison demonstrates that survival outcomes varied across "
    "both class and gender categories."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "A grouped bar chart is appropriate because it allows multiple "
    "categories to be compared simultaneously while maintaining "
    "a simple visual structure."
)

# ------------------------------------------------------------
# VISUALIZATION 3
# ------------------------------------------------------------

document.add_heading(
    "7. Visualization 3 – Age Distribution",
    level=1
)

document.add_picture(
    chart3,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 3: Distribution of passenger ages according to survival status."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "The age distribution shows how passengers of different ages "
    "were represented among survivors and non-survivors. "
    "The density curves provide an additional view of where ages "
    "were concentrated."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "A histogram combined with density estimation is useful for "
    "understanding the shape and concentration of a numerical variable."
)

# ------------------------------------------------------------
# VISUALIZATION 4
# ------------------------------------------------------------

document.add_heading(
    "8. Visualization 4 – Survival by Age Group",
    level=1
)

document.add_picture(
    chart4,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 4: Survival rates for different age categories."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "Grouping passengers into age categories makes the analysis "
    "easier to communicate. It provides a broad comparison of "
    "survival outcomes across children, teenagers, young adults, "
    "adults and seniors."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "Age groups convert a continuous numerical variable into "
    "meaningful categories, making the results easier for a "
    "general audience to interpret."
)

# ------------------------------------------------------------
# VISUALIZATION 5
# ------------------------------------------------------------

document.add_heading(
    "9. Visualization 5 – Age, Fare and Survival",
    level=1
)

document.add_picture(
    chart5,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 5: Relationship between passenger age, ticket fare, "
    "passenger class and survival."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "The scatter plot allows several variables to be viewed together. "
    "Age is shown on the horizontal axis and fare on the vertical axis. "
    "Survival status is represented using different categories, while "
    "point size provides information about passenger class."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "A scatter plot is effective for discovering relationships, "
    "clusters and unusual observations between numerical variables."
)

# ------------------------------------------------------------
# VISUALIZATION 6
# ------------------------------------------------------------

document.add_heading(
    "10. Visualization 6 – Correlation Heatmap",
    level=1
)

document.add_picture(
    chart6,
    width=Inches(6.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 6: Correlation matrix showing relationships among "
    "selected numerical variables."
)

document.add_heading("Interpretation", level=2)

document.add_paragraph(
    "The heatmap summarizes pairwise relationships between numerical "
    "variables. Stronger positive or negative relationships are easier "
    "to identify through the intensity of the displayed values."
)

document.add_heading("Why this visualization was selected", level=2)

document.add_paragraph(
    "A correlation heatmap provides a compact overview of several "
    "relationships at once and can help analysts identify variables "
    "that deserve additional investigation."
)

# ------------------------------------------------------------
# OVERALL SURVIVAL
# ------------------------------------------------------------

document.add_heading(
    "11. Overall Survival Distribution",
    level=1
)

document.add_picture(
    chart7,
    width=Inches(5.5)
)

document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

document.add_paragraph(
    "Figure 7: Overall proportion of passengers who survived and "
    "did not survive."
)

document.add_paragraph(
    f"The overall survival rate in the analyzed dataset is "
    f"{survival_rate:.2f}%. The chart provides a simple starting "
    f"point before examining more detailed factors."
)

# ------------------------------------------------------------
# DATA STORY
# ------------------------------------------------------------

document.add_heading(
    "12. Data Story – From Overview to Explanation",
    level=1
)

story = [
    (
        "Step 1 – Establishing the baseline",
        "The overall survival visualization establishes the basic "
        "proportion of survivors and non-survivors."
    ),
    (
        "Step 2 – Passenger class",
        "The analysis then examines whether passenger class is "
        "associated with different survival rates."
    ),
    (
        "Step 3 – Gender",
        "Adding gender reveals another dimension of variation "
        "within passenger classes."
    ),
    (
        "Step 4 – Age",
        "Age distributions and age groups provide insight into "
        "how survival outcomes varied across different ages."
    ),
    (
        "Step 5 – Fare and relationships",
        "The scatter plot combines multiple dimensions and helps "
        "identify patterns involving fare, age and survival."
    ),
    (
        "Step 6 – Overall relationships",
        "The correlation heatmap provides a numerical summary of "
        "relationships between important variables."
    )
]

for heading, description in story:
    document.add_heading(heading, level=2)
    document.add_paragraph(description)

# ------------------------------------------------------------
# BUSINESS IMPLICATIONS
# ------------------------------------------------------------

document.add_heading(
    "13. Potential Business and Scientific Implications",
    level=1
)

document.add_paragraph(
    "Although the Titanic dataset is historical, the analytical "
    "approach can be applied to modern business and scientific "
    "problems."
)

implications = [
    "Customer segmentation: Organizations can analyze outcomes across different customer groups.",
    "Risk analysis: Visualization can help identify groups with different levels of observed risk.",
    "Resource allocation: Data patterns can support decisions about where resources may be needed.",
    "Safety analysis: Historical incident data can be studied to identify factors associated with outcomes.",
    "Scientific research: Visualization helps researchers identify patterns and relationships before statistical modeling.",
    "Communication: Well-designed charts make analytical results easier for non-technical stakeholders to understand."
]

for item in implications:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ------------------------------------------------------------
# LIMITATIONS
# ------------------------------------------------------------

document.add_heading(
    "14. Limitations",
    level=1
)

limitations = [
    "The Titanic dataset represents a historical event and should not automatically be generalized to modern populations.",
    "Correlation between variables does not prove causation.",
    "Some passenger information contains missing values.",
    "The analysis is primarily descriptive rather than causal.",
    "Grouping continuous variables into categories can hide some individual-level variation.",
    "The dataset contains historical social and economic conditions that differ from modern circumstances."
]

for item in limitations:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ------------------------------------------------------------
# CONCLUSION
# ------------------------------------------------------------

document.add_heading(
    "15. Conclusion",
    level=1
)

document.add_paragraph(
    "This project demonstrates how Python visualization libraries "
    "can transform a dataset into a coherent data story. Multiple "
    "visualization techniques were used to investigate survival "
    "patterns across passenger class, gender, age and fare."
)

document.add_paragraph(
    "The analysis demonstrates that visualization is more than simply "
    "creating attractive charts. Effective visualization requires "
    "selecting an appropriate chart type, providing context, "
    "highlighting important patterns and communicating findings "
    "clearly to the intended audience."
)

# ------------------------------------------------------------
# TOOLS USED
# ------------------------------------------------------------

document.add_heading(
    "16. Tools and Technologies",
    level=1
)

tools_used = [
    "Python",
    "Pandas – Data manipulation and analysis",
    "NumPy – Numerical operations",
    "Matplotlib – Data visualization",
    "Seaborn – Statistical visualization",
    "python-docx – Automated report generation"
]

for item in tools_used:
    document.add_paragraph(
        item,
        style="List Bullet"
    )

# ------------------------------------------------------------
# SAVE DOCUMENT
# ------------------------------------------------------------

doc_path = os.path.join(
    OUTPUT_DIR,
    "Week_2_Titanic_Data_Storytelling_Report.docx"
)

document.save(doc_path)

print("\n==========================================")
print("PROJECT COMPLETED SUCCESSFULLY")
print("==========================================")

print("\nOutput folder:")
print(OUTPUT_DIR)

print("\nReport:")
print(doc_path)

print("\nCharts generated:")
print(chart1)
print(chart2)
print(chart3)
print(chart4)
print(chart5)
print(chart6)
print(chart7)